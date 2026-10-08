#!/usr/bin/env python3
"""辩论备赛资料生成器：Excel 资料表生成 + Markdown 转 PDF。

用法:
  generate_report.py xlsx <输出.xlsx> <数据.json>
  generate_report.py pdf  <输入.md> <输出.pdf>

xlsx 数据 JSON 结构:
{
  "sheets": [
    {
      "name":   "sheet名(≤31字符, 不含[]:*?/\\)",
      "title":  "第2行大标题(14pt加粗)",
      "note":   "第3行灰色小字注释(可省略)",
      "headers": ["编号","类别","内容","关键细节","来源"],
      "rows": [
        ["1", "统计数据", "...", "...", "[1]"],
        {"divider": "—— 当下环境 ——"},        // 灰色分隔行
        ["2", "...", "...", "...", "[2]"],
        {"summary": "小结：关键数字呈上升趋势"}  // 小结/洞察行(加粗浅灰底)
      ],
      "link_col": 4,          // URL 超链接列下标(0-based, 可省略)
      "widths": [8, 10, 40, 40, 10]  // 列宽覆盖(可省略)
    }
  ]
}
"""

import sys
import os
import json
import math
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- 样式常量 ----------
FONT_CN = "微软雅黑"   # Windows 自带黑体类，中文正文
FONT_NUM = "Cambria"   # 衬线体，编号/数字
HEADER_BG = "333333"
ALT_BG = "F9F9F9"
DIVIDER_BG = "E7E6E6"
SUMMARY_BG = "F2F2F2"
BORDER_COLOR = "DDDDDD"
LINK_COLOR = "0563C1"
GREY_NOTE = "808080"

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


def _font(size=10, bold=False, color=None, underline=None, name=FONT_CN):
    return Font(name=name, size=size, bold=bold, color=color, underline=underline)


def _col_width(header):
    """按表头名启发式分配列宽：编号窄、内容宽。"""
    if "编号" in header:
        return 8
    if "来源" in header:
        return 10
    if "URL" in header or "链接" in header:
        return 45
    if any(k in header for k in ("内容", "表述", "依据", "细节", "事实", "维度",
                                 "数据", "应对", "反驳", "核心", "起点", "现状",
                                 "结果", "对象", "观点", "类别", "人物")):
        return 38
    return 22


def _estimate_height(row_vals, widths):
    """按内容长度估算行高。"""
    max_lines = 1
    for v, w in zip(row_vals, widths):
        if v is None:
            continue
        text = str(v)
        per_line = max(4, int(w / 2))  # 中文字符约为列宽单位的 1/2
        lines = 0
        for seg in text.split("\n"):
            lines += max(1, math.ceil(len(seg) / per_line))
        max_lines = max(max_lines, lines)
    return max(18, max_lines * 14.5 + 4)


def _build_sheet(wb, spec):
    """按格式规范构建一个 sheet。"""
    name = spec["name"][:31]
    for ch in "[]:*?/\\":
        name = name.replace(ch, "_")
    ws = wb.create_sheet(title=name)
    ws.sheet_view.showGridLines = False          # 隐藏网格线
    ws.column_dimensions["A"].width = 3          # A列留白

    headers = spec.get("headers", [])
    ncols = len(headers)
    last_col = 1 + ncols                          # A=1留白, 数据从B(2)开始
    widths = spec.get("widths") or [_col_width(h) for h in headers]
    for j, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(2 + j)].width = w

    # 第2行：标题
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    c = ws.cell(row=2, column=2, value=spec.get("title", name))
    c.font = _font(size=14, bold=True)
    ws.row_dimensions[2].height = 26

    # 第3行：灰色小字注释
    if spec.get("note"):
        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last_col)
        c = ws.cell(row=3, column=2, value=spec["note"])
        c.font = _font(size=9, color=GREY_NOTE)
        ws.row_dimensions[3].height = 16

    # 第4行：表头（深灰底 + 白字加粗）
    for j, h in enumerate(headers):
        c = ws.cell(row=4, column=2 + j, value=h)
        c.font = _font(size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER
    ws.row_dimensions[4].height = 22

    # 冻结表头（前4行）
    ws.freeze_panes = "A5"

    # 第5行起：数据
    link_col = spec.get("link_col")
    r = 5
    for item in spec.get("rows", []):
        if isinstance(item, dict) and "divider" in item:
            # 灰色分隔行：合并整行、浅灰底、灰色文字
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)
            c = ws.cell(row=r, column=2, value=item["divider"])
            c.font = _font(size=10, bold=True, color="595959")
            c.fill = PatternFill("solid", fgColor=DIVIDER_BG)
            c.alignment = Alignment(horizontal="left", vertical="center")
            for j in range(ncols):
                ws.cell(row=r, column=2 + j).border = THIN_BORDER
            ws.row_dimensions[r].height = 20
            r += 1
            continue
        if isinstance(item, dict) and "summary" in item:
            # 小结/洞察行：合并整行、加粗、浅灰底
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)
            c = ws.cell(row=r, column=2, value=item["summary"])
            c.font = _font(size=10, bold=True)
            c.fill = PatternFill("solid", fgColor=SUMMARY_BG)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            for j in range(ncols):
                ws.cell(row=r, column=2 + j).border = THIN_BORDER
            ws.row_dimensions[r].height = _estimate_height([item["summary"]], [sum(widths)])
            r += 1
            continue

        row = list(item) + [None] * (ncols - len(item))
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=2 + j, value=v)
            is_num_col = "编号" in headers[j] or isinstance(v, (int, float))
            if link_col is not None and j == link_col and v:
                # 超链接列：蓝色下划线 + 可点击
                c.hyperlink = str(v)
                c.font = _font(size=10, color=LINK_COLOR, underline="single")
            else:
                c.font = _font(size=10, name=FONT_NUM) if is_num_col else _font(size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if is_num_col else "left")
            c.border = THIN_BORDER
            if (r - 5) % 2 == 1:  # 隔行交替：F9F9F9
                c.fill = PatternFill("solid", fgColor=ALT_BG)
        ws.row_dimensions[r].height = _estimate_height(row, widths)
        r += 1


def build_xlsx(data, out_path):
    """从 JSON 数据生成 xlsx。"""
    wb = Workbook()
    wb.remove(wb.active)
    for spec in data.get("sheets", []):
        _build_sheet(wb, spec)
    wb.save(out_path)
    print(f"XLSX: {out_path}")


# ---------- PDF ----------

CSS = """
@page { size: A4; margin: 2.2cm 2cm; }
body { font-family: "Noto Sans CJK SC", sans-serif; font-size: 10.5pt; line-height: 1.65; color: #1a1a1a; }
h1 { font-size: 20pt; text-align: center; margin: 0 0 6pt; }
h2 { font-size: 14pt; margin: 14pt 0 6pt; padding-bottom: 3pt; border-bottom: 2px solid #333333; }
h3 { font-size: 12pt; margin: 10pt 0 4pt; }
p { margin: 4pt 0; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 10pt; font-size: 9.5pt; }
th { background: #333333; color: #ffffff; font-weight: bold; padding: 4pt 6pt; border: 1px solid #cccccc; text-align: left; }
td { padding: 4pt 6pt; border: 1px solid #dddddd; vertical-align: top; }
tr:nth-child(even) td { background: #f9f9f9; }
a { color: #0563c1; text-decoration: underline; }
blockquote { margin: 6pt 0; padding: 4pt 10pt; border-left: 3px solid #cccccc; color: #555555; background: #fafafa; }
code { background: #f2f2f2; padding: 0 3pt; }
hr { border: none; border-top: 1px solid #cccccc; margin: 10pt 0; }
"""


def md_to_pdf(md_path, pdf_path):
    """pandoc md→html（内嵌 CSS）→ weasyprint html→pdf。"""
    with tempfile.TemporaryDirectory() as td:
        # CSS 必须内联进 HTML：Windows 下 pandoc -c 会把绝对路径写成
        # 反斜杠/百分号编码的伪 URI，weasyprint 解析不了（Not an absolute URI）。
        header_path = os.path.join(td, "header.html")
        with open(header_path, "w", encoding="utf-8") as f:
            f.write("<style>\n" + CSS + "\n</style>\n")
        html_path = os.path.join(td, "report.html")
        subprocess.run(
            ["pandoc", "-s", md_path, "-o", html_path,
             "--metadata", "title=辩题调查报告", "-H", header_path],
            check=True,
        )
        subprocess.run(["weasyprint", html_path, pdf_path], check=True)
    print(f"PDF:  {pdf_path}")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    cmd = sys.argv[1]
    if cmd == "xlsx" and len(sys.argv) == 4:
        with open(sys.argv[3], encoding="utf-8") as f:
            data = json.load(f)
        build_xlsx(data, sys.argv[2])
    elif cmd == "pdf" and len(sys.argv) == 4:
        md_to_pdf(sys.argv[2], sys.argv[3])
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
